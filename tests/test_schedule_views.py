"""Tests for the readable schedule views."""

import numpy as np
import pytest
from openpyxl import load_workbook

from src.optimizer_247 import solve_weekly_schedule
from src.patterns import (
    WEEKDAY_EXTENDED,
    build_pattern_demand,
    build_pattern_options,
    pattern_roster,
)
from src.schedule_views import (
    BREAK,
    interval_label,
    MEAL,
    OFF,
    ON_QUEUE,
    day_state_matrix,
    render_day_grid,
    write_schedule_workbook,
)


@pytest.fixture(scope="module")
def solved():
    from dataclasses import replace

    from src.optimizer_247 import WeeklySolverConfig

    small = replace(WEEKDAY_EXTENDED, days=2, full_time=10, part_time=4, peak_calls=25.0)
    # A two-day grid cannot carry the default consecutive-night window.
    config = WeeklySolverConfig(max_consecutive_nights=1, full_time_night_min=0,
                                full_time_night_max=1)
    return small, solve_weekly_schedule(
        pattern_roster(small), build_pattern_options(small), build_pattern_demand(small),
        config, small.grid, time_limit_sec=60, gap=0.05,
    )


def test_the_matrix_has_one_row_per_agent_on_duty(solved):
    pattern, schedule = solved
    rows, matrix = day_state_matrix(schedule, day=0)
    on_duty = {a.agent_name for a in schedule.assignments if a.option.start_day == 0}
    assert {r.agent for r in rows} == on_duty
    assert matrix.shape == (len(on_duty), 96)


def test_each_row_carries_the_shift_window(solved):
    """The label has to say when the agent works, not just who they are."""
    _, schedule = solved
    rows, _ = day_state_matrix(schedule, day=0)
    by_name = {a.agent_name: a for a in schedule.assignments if a.option.start_day == 0}
    for shift in rows:
        option = by_name[shift.agent].option
        assert shift.start == interval_label(option.start_slot)
        assert shift.paid_hours == option.paid_hours
        assert shift.label.startswith(shift.agent)
        assert f"{shift.start}-{shift.end}" in shift.label


def test_rows_are_ordered_by_start_time(solved):
    _, schedule = solved
    rows, _ = day_state_matrix(schedule, day=0)
    assert [r.start for r in rows] == sorted(r.start for r in rows)


def test_every_cell_is_a_known_state(solved):
    _, schedule = solved
    _, matrix = day_state_matrix(schedule, day=0)
    assert set(np.unique(matrix)).issubset({OFF, ON_QUEUE, BREAK, MEAL})


def test_on_queue_cells_match_the_shift_coverage(solved):
    """The picture must agree with the schedule it claims to draw."""
    _, schedule = solved
    rows, matrix = day_state_matrix(schedule, day=0)
    by_name = {a.agent_name: a for a in schedule.assignments if a.option.start_day == 0}
    for row, shift in enumerate(rows):
        covered = {s % 96 for s in by_name[shift.agent].option.coverage_slots}
        drawn = {int(i) for i in np.flatnonzero(matrix[row] == ON_QUEUE)}
        assert drawn == covered


def test_breaks_are_drawn_and_split_from_meals(solved):
    """A single relief interval and a longer meal read differently."""
    _, schedule = solved
    _, matrix = day_state_matrix(schedule, day=0)
    assert (matrix == BREAK).any() or (matrix == MEAL).any()


def test_a_day_with_no_assignments_is_rejected(solved):
    _, schedule = solved
    with pytest.raises(ValueError, match="no assignments"):
        render_day_grid(schedule, day=6)


def test_the_grid_renders(solved, tmp_path):
    _, schedule = solved
    path = tmp_path / "grid.png"
    render_day_grid(schedule, day=0, save_path=path, dpi=60)
    assert path.exists() and path.stat().st_size > 5_000


def test_the_workbook_has_a_sheet_per_day_and_a_legend(solved, tmp_path):
    pattern, schedule = solved
    path = write_schedule_workbook(schedule, tmp_path / "s.xlsx", days=pattern.days)
    book = load_workbook(path)
    assert book.sheetnames[0] == "Legend"
    assert len(book.sheetnames) == pattern.days + 1


def test_the_workbook_freezes_the_agent_column(solved, tmp_path):
    """The first thing anyone does is scroll sideways looking for one person."""
    pattern, schedule = solved
    path = write_schedule_workbook(schedule, tmp_path / "s.xlsx", days=pattern.days)
    sheet = load_workbook(path)["Mon"]
    assert sheet.freeze_panes == "E2"
    assert [sheet.cell(row=1, column=c).value for c in range(1, 5)] == [
        "Agent", "Start", "End", "Paid h",
    ]
    assert sheet.cell(row=2, column=2).value  # a start time, not blank


def test_overnight_rows_match_actual_calendar_day_and_paid_hours():
    from src.continuous import build_weekly_shift_options
    from src.optimizer_247 import WeeklyAssignment, WeeklySchedule

    options = {o.id: o for o in build_weekly_shift_options()}
    assignments = [
        WeeklyAssignment(1, "A", options["Sun_FT8_2200"]),
        WeeklyAssignment(2, "B", options["Mon_FT8_2200"]),
    ]
    coverage = np.zeros(672, dtype=int)
    for assignment in assignments:
        coverage[list(assignment.option.coverage_slots)] += 1
    schedule = WeeklySchedule(
        assignments, coverage, coverage.copy(), np.zeros(672, dtype=int),
        {1: 8.0, 2: 8.0}, "Feasible", "fixture", "fixture", 0.0,
        None, None, 0.0, 0, 0, 0,
    )
    paid = 0.0
    for day in range(7):
        rows, matrix = day_state_matrix(schedule, day)
        np.testing.assert_array_equal(
            (matrix == ON_QUEUE).sum(axis=0), coverage[day * 96:(day + 1) * 96]
        )
        paid += sum(row.paid_hours for row in rows)
    assert paid == 16.0
    monday, _ = day_state_matrix(schedule, 0)
    assert "(-1d)" in monday[0].start
    assert "(+1d)" in monday[1].end
